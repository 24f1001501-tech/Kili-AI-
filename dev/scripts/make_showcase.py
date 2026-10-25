"""SayBetter AI — feature showcase (Word) with screenshots."""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT, SHOTS = sys.argv[1], Path(sys.argv[2])
INK = RGBColor(0x2B, 0x2D, 0x42)
CORAL = RGBColor(0xC8, 0x32, 0x3F)
GRAPE = RGBColor(0x6A, 0x2F, 0xB8)
GREY = RGBColor(0x4F, 0x53, 0x70)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin = sec.bottom_margin = Cm(1.8)
CONTENT_W = 17.0

st = doc.styles
st["Normal"].font.name = "Calibri"
st["Normal"].font.size = Pt(10.5)
st["Normal"].font.color.rgb = INK
st["Normal"].element.rPr.rFonts.set(qn("w:cs"), "Nirmala UI")
st["Normal"].paragraph_format.space_after = Pt(4)
for name, size, color in (("Title", 30, INK), ("Heading 1", 18, CORAL), ("Heading 2", 13, GRAPE)):
    s = st[name]
    s.font.name = "Calibri"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = color
    s.element.rPr.rFonts.set(qn("w:cs"), "Nirmala UI")
st["Heading 1"].paragraph_format.space_before = Pt(14)
st["Heading 2"].paragraph_format.space_before = Pt(10)
st["Heading 2"].paragraph_format.keep_with_next = True


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def para(text="", bold_prefix=None, italic=False, color=None, size=None, align=None, bold=False, keep=False):
    p = doc.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    r = p.add_run(text)
    r.italic, r.bold = italic, bold
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    p.paragraph_format.keep_with_next = keep
    return p


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            p.add_run(it[0]).bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)


def table(headers, rows, widths, font=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        shade(c, "FFD93D")
        c.paragraphs[0].paragraph_format.keep_with_next = True
    for k, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(str(val))
            if k % 2:
                shade(cells[i], "F7F5FB")
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
            for p in row.cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(font)
    doc.add_paragraph()
    return t


def image(name, width_cm=CONTENT_W, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = bool(caption)
    p.add_run().add_picture(str(SHOTS / f"{name}.png"), width=Cm(width_cm))
    if caption:
        c = para(caption, italic=True, color=GREY, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        c.paragraph_format.space_after = Pt(8)


def feature(title, shot, width, shows, demo=None, caption=None):
    doc.add_heading(title, level=2)
    image(shot, width, caption)
    bullets(shows)
    if demo:
        para(demo, bold_prefix="Demo tip: ", color=GREY, size=9.5)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------- title ----------------
for _ in range(3):
    doc.add_paragraph()
doc.add_paragraph("SayBetter AI", style="Title")
para("Feature showcase", color=CORAL, size=20, bold=True)
para("A speech-practice coach for Tamil and Hindi — every screen, what it does, and how to demo it", color=GREY, size=12)
doc.add_paragraph()
image("01_landing", 14.5)
para("MVP v0.2 · candy-pop redesign · 24 October 2026", color=GREY, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
page_break()

# ---------------- at a glance ----------------
doc.add_heading("At a glance", level=1)
para("A learner says a word. SayBetter checks the one sound that matters, tells them which sound it actually heard "
     "(“We heard स instead of श”), shows the tongue and lips for both sounds side by side, gives one specific tip, and "
     "moves their plan up or back based on saved results. It supports practice; it does not diagnose or replace a "
     "speech-language therapist.")
table(["What", "How"], [
    ["Languages and sounds", "Tamil ழ ள ர · Hindi र श ख — each with 5 levels from single sounds to everyday sentences"],
    ["Three checks per attempt", "Sarvam speech-to-text · a phoneme model on the computer · comparison with 6 Sarvam voices saying the word correctly and with each mix-up"],
    ["Honest confidence", "Each check counts in proportion to its measured accuracy for that sound; below 55% the app says “not sure”"],
    ["Adaptive plan", "3 of last 4 on target → level up · 3 of 4 mixed up → level down · stuck at level 1 → suggests a therapist"],
    ["Privacy", "Recordings deleted after analysis unless the learner opts in; one-tap delete of all data"],
], [4.2, 12.8])

# ---------------- design ----------------
doc.add_heading("Design system", level=1)
para("Candy-pop colours on white, cut-out “sticker” surfaces with thick ink outlines and hard shadows, and a parrot "
     "mascot. Colours are used as fills; text always uses ink or a deeper shade so it stays readable.")
t = doc.add_table(rows=2, cols=6)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
swatches = [("Coral", "FF6B6B", "Primary buttons"), ("Sunshine", "FFD93D", "Highlights, active nav"),
            ("Sky", "4D96FF", "On target, Kili"), ("Grape", "9B5DE5", "Links, roof of mouth"),
            ("Ink", "2B2D42", "Text, outlines"), ("White", "FFFFFF", "Page background")]
for i, (name, hexv, role) in enumerate(swatches):
    top, bottom = t.rows[0].cells[i], t.rows[1].cells[i]
    shade(top, hexv)
    top.width = bottom.width = Cm(2.8)
    top.text = ""
    top.paragraphs[0].add_run(" ").font.size = Pt(26)
    bottom.text = ""
    bp = bottom.paragraphs[0]
    bp.add_run(f"{name}\n").bold = True
    bp.add_run(f"#{hexv}\n{role}")
    for r in bp.runs:
        r.font.size = Pt(8.5)
doc.add_paragraph()
table(["Element", "Choice", "Why"], [
    ["Headings", "Baloo 2 (ExtraBold)", "Rounded and friendly; made by an Indian type foundry with matching Devanagari, so Hindi words share the style"],
    ["Tamil words", "Baloo Thambi 2", "The Tamil sibling of Baloo — ழ, ள, ர look as bubbly as the English text"],
    ["Body text", "DM Sans", "Clean and readable at small sizes"],
    ["Mascot", "Kili the parrot", "“Kili” is Tamil for parrot — and parrots copy speech"],
], [3.0, 4.0, 10.0])
table(["Kili’s mood", "When it appears"], [
    ["Idle (blinks)", "Waiting for you to start"],
    ["Listening (head tilt + sound waves)", "While the microphone is recording"],
    ["Thinking (three dots)", "While the attempt is analysed, or when the result is “not sure”"],
    ["Happy (hops, sparkles)", "The target sound was on target"],
    ["Encouraging (wing up)", "A mix-up was heard — never a sad face"],
], [6.0, 11.0])
para("Motion is subtle and always means something: a live waveform while recording, Kili’s reactions, a gentle "
     "page fade, the confidence bar filling, and progress bars growing in. All of it switches off for people who "
     "set “reduce motion” on their device.", size=10)
page_break()

# ---------------- walkthrough ----------------
doc.add_heading("Feature walkthrough", level=1)
para("The screens below follow a real session: a Tamil learner (Priya) and a Hindi learner (Arjun). Their speech was "
     "simulated with Sarvam’s voices, including deliberate mix-ups, and analysed by the real system.", italic=True, color=GREY)

feature("1. Landing page", "01_landing", 15.0, [
    ("Headline and sticker letters: ", "the six practice sounds in Tamil and Hindi, set in the brand typefaces."),
    ("Live mouth diagram: ", "the tongue curls back for ழ while air flows over it — the product’s core idea on the first screen."),
    ("Kili introduces itself; ", "privacy and “not a diagnosis” notes sit right below the fold."),
], "Start here to explain the problem: most speech apps only cover English sounds.")

feature("2. Onboarding", "02_onboarding", 13.0, [
    "Name, language (Tamil or Hindi), an optional goal, and feedback style (brief, step by step or encouraging).",
    "Recordings are not kept unless the learner opts in; microphone access is asked for only when they press record.",
])
page_break()

feature("3. Two-minute check", "03_baseline", 14.0, [
    "Six short words, two per sound, with the target sound highlighted, romanised and translated.",
    "“Hear it” plays a native reference voice from Sarvam; the recorder accepts the microphone or a WAV upload.",
])
feature("4. Plan ready", "04_plan_ready", 12.0, [
    "Sounds are ranked by how they came out: here ழ needs focus, ள is next, and ர sounded clear.",
    "Starting levels are set from the results; the note reminds learners this is not an assessment.",
])
page_break()

feature("5. Practice room", "05_practice", 14.0, [
    ("Word card: ", "Kili next to a speech bubble with the word, its sound highlighted."),
    ("How-to panel: ", "an animated side view of the mouth plus four plain-language steps."),
    ("Level badge: ", "shows the current level (here level 2, short words)."),
])
feature("6. Recording", "06_recording", 7.2, [
    "The waveform moves with the learner’s voice, so they can see the microphone is hearing them.",
    "Kili tilts its head and shows sound waves while listening.",
], "In Chrome or Edge, record live. In the app’s built-in browser panel the microphone is blocked, so use WAV upload there.")
page_break()

feature("7. Feedback — Hindi mix-up (the key moment)", "13_feedback_hindi_mixup", 15.5, [
    ("Verdict: ", "“We heard स instead of श”, with high confidence because all three checks agreed."),
    ("Evidence: ", "Sarvam wrote सा; the phoneme model heard /s aː/; the closest Sarvam voice was the mix-up."),
    ("One tip: ", "pull the tongue back, round the lips, blow “shhh”."),
    ("Side-by-side mouths: ", "the target sound and the sound that was heard."),
], "Open “Why this confidence?” to show each check’s measured accuracy for this sound.")
page_break()
feature("8. Feedback — correct attempt", "14_feedback_hindi_correct", 15.5, [
    "Kili hops with sparkles and a check mark appears; “Continue” becomes the main button.",
    "Three on-target attempts out of the last four move the learner up a level automatically.",
])
feature("9. Feedback — honest “not sure” (Tamil ழ)", "07_feedback_tamil", 15.5, [
    "For ழ, no check has yet proven reliable enough on its own, so the app says “not sure” instead of guessing.",
    "It still shows what each check heard and gives a step to focus on. On the retry with a clear ழ, the app marked it on target.",
], "Use this screen to show judges the system is honest about its limits.")
page_break()

feature("10. Sound guide", "09_sound_guide", 14.0, [
    "How to make the sound, a large animated diagram, and every common mix-up with its own comparison and tip.",
    "All five levels with their practice words; one button sets this sound as the focus.",
])
page_break()
feature("11. Dashboard", "10_dashboard", 15.0, [
    ("Current focus: ", "a tilted sticker letter, a stepping-stone path for the 5 levels, and the next word."),
    ("Your sounds: ", "status for every sound, with a button to switch focus."),
    ("Recent attempts: ", "coloured dots for on target, mixed up and unsure, with a legend."),
    ("Status chips: ", "show that Sarvam and the phoneme model are connected."),
])
page_break()
feature("12. Progress", "11_progress", 15.0, [
    "Sessions, practice days and recordings, then one card per sound with its level path, on-target share and a daily bar chart that grows in.",
    "Labelled as an app estimate of practice performance, not clinical improvement.",
])
feature("13. Settings and measured accuracy", "12_settings_accuracy", 11.0, [
    "Profile, feedback style and recording consent; delete all data in one step.",
    "The measured-accuracy table shows how often the final verdict and each check were right for each sound.",
])
page_break()

doc.add_heading("14. On a phone", level=2)
t = doc.add_table(rows=2, cols=3)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (shot, cap) in enumerate((("17_mobile_landing", "Landing"), ("15_mobile_practice", "Practice"), ("16_mobile_dashboard", "Dashboard"))):
    c = t.rows[0].cells[i]
    c.width = Cm(5.6)
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraphs[0].add_run().add_picture(str(SHOTS / f"{shot}.png"), width=Cm(5.2))
    cc = t.rows[1].cells[i]
    cc.text = ""
    r = cc.paragraphs[0].add_run(cap)
    r.italic, r.font.size = True, Pt(9)
    cc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()
bullets(["Every screen fits a 390-px phone with no sideways scrolling; the navigation becomes one compact row.",
         "Buttons are at least 48 px tall for easy tapping."])
page_break()

# ---------------- accuracy ----------------
doc.add_heading("Measured accuracy", level=1)
para("Measured on 1,218 clips: every practice word at levels 1–3, said correctly and with each common mix-up, by 6 Sarvam "
     "voices. The voice comparison never compares a clip against its own voice.")
table(["Sound", "Final verdict right", "Wrong", "“Not sure”", "Right when it made a call"], [
    ["श (sha)", "95%", "5%", "0%", "95%"],
    ["र (ra, Hindi)", "80%", "20%", "0%", "80%"],
    ["ர (ra, Tamil)", "75%", "25%", "1%", "75%"],
    ["ख (kha)", "74%", "23%", "2%", "76%"],
    ["ழ (zha)", "30%", "10%", "60%", "76%"],
    ["ள (La)", "0%", "0%", "100%", "—"],
], [3.4, 3.4, 2.2, 2.6, 5.4])
bullets([
    ("Best case: ", "these are synthetic voices, so real learners will score lower; the accuracy needs checking with native speakers."),
    ("Optimistic for the final verdict: ", "the confidence weights were learned from the same clips they are tested on."),
    ("Tamil ழ and ள are the hardest: ", "the app mostly says “not sure” for them — the honest behaviour — until real recordings improve the models."),
])

# ---------------- demo script ----------------
doc.add_heading("Three-minute demo script", level=1)
table(["Time", "Show", "Say"], [
    ["0:00", "Landing page", "“Half of India speaks Hindi or Tamil, but speech apps only cover English sounds.”"],
    ["0:30", "Two-minute check → plan", "“Six words find the sound to work on first, like a therapist’s assessment.”"],
    ["1:00", "Practice: record शा, then सा", "“It names the exact sound it heard and shows both mouths side by side.”"],
    ["1:45", "Why this confidence?", "“Three checks, weighted by measured accuracy. If it isn’t sure, it says so.”"],
    ["2:15", "Dashboard and progress", "“The plan moves up or back on evidence, and parents see real practice data.”"],
    ["2:40", "Settings → accuracy", "“We publish our accuracy, and recordings are deleted by default.”"],
], [1.6, 5.0, 10.4])

doc.add_heading("Running it", level=1)
para("powershell -ExecutionPolicy Bypass -File saybetter\\run.ps1", size=9.5)
para("Then open http://localhost:8000 in Chrome or Edge. The Sarvam key goes in saybetter/backend/.env.", size=10)

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run("SayBetter AI — Feature showcase   ·   page ")
run.font.size = Pt(8)
run.font.color.rgb = GREY
for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
    r = fp.add_run()
    r.font.size = Pt(8)
    if tag:
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), tag)
        r._r.append(fc)
    else:
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = text
        r._r.append(it)

doc.save(OUT)
print("saved", OUT)
